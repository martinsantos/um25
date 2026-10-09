from pathlib import Path
import json, uuid, zipfile, shutil
from lxml import etree
from docx import Document
from docx.shared import Pt, Mm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

import os, tempfile
ROOT=Path(os.environ.get("UM_TEMPLATE_WORK",Path(tempfile.gettempdir())/"um-sans-template-build"))
REPO=Path(os.environ.get("UM_TEMPLATE_REPO",Path(__file__).resolve().parents[3]))
OUT=ROOT/'entrega'
BASE=REPO/'public/fonts/um-sans/v2.0.0'
FONT='UM Sans 2'
SIZES=[16,13,12]
TABS=[12,17,23]
NS={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}

def el(tag,**attrs):
    x=OxmlElement('w:'+tag)
    for k,v in attrs.items():x.set(qn('w:'+k),str(v))
    return x

def face(style,size,bold=False,color='000000'):
    style.font.name=FONT;style.font.size=Pt(size);style.font.bold=bold;style.font.italic=False
    style.font.color.rgb=RGBColor.from_string(color)
    rpr=style.element.get_or_add_rPr()
    f=rpr.rFonts
    for k in list(f.attrib):del f.attrib[k]
    for k in ['ascii','hAnsi','cs','eastAsia']:f.set(qn('w:'+k),FONT)
    for ch in list(rpr):
        if ch.tag in [qn('w:spacing'),qn('w:iCs')]:rpr.remove(ch)
    for ch in list(style.element.get_or_add_pPr()):
        if ch.tag in [qn('w:pBdr'),qn('w:numPr')]:style.element.get_or_add_pPr().remove(ch)

def field(p,code):
    f=el('fldSimple',instr=code)
    r=el('r');r.append(el('t'));r[-1].text='1';f.append(r);p._p.append(f)

def numbering(d):
    n=d.part.numbering_part.element
    a=el('abstractNum',abstractNumId=42);a.append(el('multiLevelType',val='multilevel'))
    for i,(size,tab) in enumerate(zip(SIZES,TABS)):
        lv=el('lvl',ilvl=i);lv.append(el('start',val=1));lv.append(el('numFmt',val='decimal'))
        lv.append(el('pStyle',val=f'Heading{i+1}'))
        lv.append(el('lvlText',val='.'.join('%'+str(j+1) for j in range(i+1))))
        lv.append(el('suff',val='tab'));lv.append(el('lvlJc',val='left'))
        if i:lv.append(el('lvlRestart',val=i))
        p=el('pPr');tabs=el('tabs');tabs.append(el('tab',val='num',pos=round(tab*1440/25.4)));p.append(tabs)
        p.append(el('ind',left=round(tab*1440/25.4),hanging=round(tab*1440/25.4)));lv.append(p)
        r=el('rPr');r.append(el('rFonts',ascii=FONT,hAnsi=FONT,cs=FONT,eastAsia=FONT));r.append(el('sz',val=size*2));r.append(el('b'));r.append(el('color',val='000000'));lv.append(r);a.append(lv)
    n.append(a);num=el('num',numId=42);num.append(el('abstractNumId',val=42));n.append(num)
    for i in range(3):
        st=d.styles[f'Heading {i+1}'];np=el('numPr');np.append(el('ilvl',val=i));np.append(el('numId',val=42));st.element.get_or_add_pPr().append(np)
        st.paragraph_format.left_indent=Mm(TABS[i]);st.paragraph_format.first_line_indent=Mm(-TABS[i])
        st.paragraph_format.tab_stops.add_tab_stop(Mm(TABS[i]))
        st.next_paragraph_style=d.styles['Normal']

def new_doc(title,subtitle,proof=False,letter=False):
    d=Document();s=d.sections[0]
    s.page_width=Mm(210);s.page_height=Mm(297)
    s.left_margin=Mm(25);s.right_margin=Mm(20);s.top_margin=Mm(31);s.bottom_margin=Mm(24)
    s.header_distance=Mm(14);s.footer_distance=Mm(12)
    for name,size,bold in [('Normal',12,False),('Title',22,True),('Subtitle',13,False),('Heading 1',16,True),('Heading 2',13,True),('Heading 3',12,True),('Caption',9,False),('Header',9,False),('Footer',9,False)]:
        st=d.styles[name];face(st,size,bold,'666666' if name in ['Caption','Footer','Subtitle'] else '000000')
        pf=st.paragraph_format;pf.line_spacing=1.2;pf.space_after=Pt(8);pf.space_before=Pt(0);pf.widow_control=True
        if name.startswith('Heading'):
            pf.keep_with_next=True;pf.keep_together=True;pf.space_before=Pt(20 if name=='Heading 1' else 12);pf.space_after=Pt(7)
        if name=='Title':pf.line_spacing=1.08;pf.keep_with_next=True;pf.space_after=Pt(12)
        if name in ['Header','Footer']:pf.space_after=Pt(0);pf.line_spacing=1;pf.tab_stops.clear_all()
    numbering(d)
    h=s.header.paragraphs[0];h.style='Header';h.alignment=WD_ALIGN_PARAGRAPH.LEFT if letter else WD_ALIGN_PARAGRAPH.RIGHT
    h.add_run().add_picture(str(ROOT/'logo.png'),width=Mm(54))
    inline=h._p.xpath('.//wp:inline')[0]
    anchor=OxmlElement('wp:anchor')
    for k,v in {'distT':'0','distB':'0','distL':'0','distR':'0','simplePos':'0','relativeHeight':'0','behindDoc':'0','locked':'1','layoutInCell':'1','allowOverlap':'1'}.items():anchor.set(k,v)
    simple=OxmlElement('wp:simplePos');simple.set('x','0');simple.set('y','0');anchor.append(simple)
    for axis,relative,value in [('H','margin','left' if letter else 'right'),('V','paragraph','0')]:
        pos=OxmlElement('wp:position'+axis);pos.set('relativeFrom',relative)
        align=OxmlElement('wp:align' if axis=='H' else 'wp:posOffset');align.text=value;pos.append(align);anchor.append(pos)
    anchor.append(inline.find(qn('wp:extent')))
    effect=OxmlElement('wp:effectExtent')
    for k in ['l','t','r','b']:effect.set(k,'0')
    anchor.append(effect);anchor.append(OxmlElement('wp:wrapNone'))
    for name in ['wp:docPr','wp:cNvGraphicFramePr','a:graphic']:anchor.append(inline.find(qn(name)))
    inline.getparent().replace(inline,anchor)
    h.paragraph_format.line_spacing=Pt(18)
    border=el('pBdr');border.append(el('bottom',val='single',sz=6,space=7,color='DC2626'));h._p.get_or_add_pPr().append(border)
    for x in h._p.xpath('.//wp:docPr'):x.set('descr','Logotipo de ULTIMA MILLA')
    f=s.footer.paragraphs[0];f.style='Footer';f.paragraph_format.tab_stops.add_tab_stop(Mm(165),WD_TAB_ALIGNMENT.RIGHT)
    f.add_run('ULTIMA MILLA S.A. · ultimamilla.com.ar\t');field(f,'PAGE');f.add_run(' / ');field(f,'NUMPAGES')
    border=el('pBdr');border.append(el('top',val='single',sz=4,space=6,color='DDDDDD'));f._p.get_or_add_pPr().append(border)
    d.settings.element.append(el('updateFields',val='true'))
    d.settings.element.append(el('embedTrueTypeFonts'))
    d.core_properties.author='ULTIMA MILLA S.A.';d.core_properties.title=title
    d.core_properties.subject='Plantilla de licitación con estilos y numeración multinivel'
    d.add_paragraph('COMUNICACIÓN INSTITUCIONAL' if letter else 'LICITACIÓN [NÚMERO]  ·  EXPEDIENTE [REFERENCIA]' if not proof else 'EJEMPLO DE MAQUETACIÓN  ·  DATOS FICTICIOS',style='Caption')
    d.add_paragraph(title,style='Title');d.add_paragraph(subtitle,style='Subtitle')
    for label,text in [('Destinatario','[Organismo o cliente]'),('Proyecto','[Nombre del proyecto]'),('Fecha','[DD/MM/AAAA]')]:
        if letter and label=='Proyecto':continue
        p=d.add_paragraph();p.paragraph_format.space_after=Pt(5)
        p.paragraph_format.tab_stops.add_tab_stop(Mm(30))
        label_run=p.add_run(label+'\t');label_run.font.size=Pt(10);label_run.font.color.rgb=RGBColor.from_string('666666')
        p.add_run(text)
    return d

def heading(d,text,level=1,newpage=False):
    p=d.add_paragraph(text,style=f'Heading {level}')
    if newpage:p.paragraph_format.page_break_before=True
    return p

def table(d,rows,widths,aligns=None):
    t=d.add_table(rows=0,cols=len(widths));t.alignment=WD_TABLE_ALIGNMENT.LEFT;t.autofit=False
    for c,w in zip(t.columns,widths):c.width=Mm(w)
    pr=t._tbl.tblPr
    borders=el('tblBorders')
    for edge in ['top','bottom','insideH']:borders.append(el(edge,val='single',sz=4,color='DDDDDD'))
    for edge in ['left','right','insideV']:borders.append(el(edge,val='nil'))
    pr.append(borders);pr.append(el('tblInd',w=120,type='dxa'))
    margins=el('tblCellMar')
    for edge,v in [('top',100),('bottom',100),('left',120),('right',120)]:margins.append(el(edge,w=v,type='dxa'))
    pr.append(margins)
    for ri,row in enumerate(rows):
        rr=t.add_row();rr._tr.get_or_add_trPr().append(el('cantSplit'))
        if ri==0:rr._tr.get_or_add_trPr().append(el('tblHeader'))
        for ci,(c,text,w) in enumerate(zip(rr.cells,row,widths)):
            c.width=Mm(w);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            c._tc.get_or_add_tcPr().append(el('shd',fill='000000' if ri==0 else ('F5F5F5' if ri%2==0 else 'FFFFFF')))
            p=c.paragraphs[0];p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1.12
            p.alignment=(aligns[ci] if aligns and ri>0 else WD_ALIGN_PARAGRAPH.LEFT)
            r=p.add_run(text);r.font.size=Pt(10);r.bold=(ri==0);r.font.color.rgb=RGBColor.from_string('FFFFFF' if ri==0 else '333333')
    p=d.add_paragraph();p.paragraph_format.space_after=Pt(0);p.paragraph_format.space_before=Pt(0);p.paragraph_format.line_spacing=Pt(3);p.add_run().font.size=Pt(3)
    return t

def embed_fonts(path):
    with zipfile.ZipFile(path) as z:files={n:z.read(n) for n in z.namelist()}
    root=etree.fromstring(files['word/fontTable.xml']);font=next((f for f in root if f.get(qn('w:name'))==FONT),None)
    if font is None:font=el('font',name=FONT);root.append(font)
    relns='http://schemas.openxmlformats.org/package/2006/relationships'
    rels=etree.Element('{'+relns+'}Relationships',nsmap={None:relns})
    for i,(kind,file) in enumerate([('Regular','Regular'),('Bold','Bold'),('Italic','RegularItalic'),('BoldItalic','BoldItalic')]):
        guid=uuid.uuid5(uuid.NAMESPACE_URL,'UMSans2-'+file);key=guid.bytes[::-1];data=bytearray((BASE/f'UMSans2-{file}.ttf').read_bytes())
        for j in range(32):data[j]^=key[j%16]
        fname=f'font{i}.odttf';files['word/fonts/'+fname]=bytes(data)
        e=el('embed'+kind,fontKey='{'+str(guid).upper()+'}');e.set(qn('r:id'),f'rIdUM{i}');font.append(e)
        etree.SubElement(rels,'{'+relns+'}Relationship',Id=f'rIdUM{i}',Type='http://schemas.openxmlformats.org/officeDocument/2006/relationships/font',Target='fonts/'+fname)
    files['word/fontTable.xml']=etree.tostring(root,xml_declaration=True,encoding='UTF-8',standalone=True)
    files['word/_rels/fontTable.xml.rels']=etree.tostring(rels,xml_declaration=True,encoding='UTF-8',standalone=True)
    types=etree.fromstring(files['[Content_Types].xml']);ct='http://schemas.openxmlformats.org/package/2006/content-types'
    etree.SubElement(types,'{'+ct+'}Default',Extension='odttf',ContentType='application/vnd.openxmlformats-officedocument.obfuscatedFont')
    files['[Content_Types].xml']=etree.tostring(types,xml_declaration=True,encoding='UTF-8',standalone=True)
    with zipfile.ZipFile(path,'w',zipfile.ZIP_DEFLATED) as z:
        for n,b in files.items():z.writestr(n,b)

def save(d,name,template=True):
    p=OUT/(name+'.docx');d.save(p);embed_fonts(p)
    if template:
        with zipfile.ZipFile(p) as src,zipfile.ZipFile(OUT/(name+'.dotx'),'w',zipfile.ZIP_DEFLATED) as z:
            for n in src.namelist():
                b=src.read(n)
                if n=='[Content_Types].xml':b=b.replace(b'wordprocessingml.document.main+xml',b'wordprocessingml.template.main+xml')
                z.writestr(n,b)
    print(p)

def full(proof=False):
    d=new_doc('Propuesta técnica y económica','[Objeto de la contratación]',proof)
    heading(d,'Presentación de la oferta');d.add_paragraph('[Identificar al oferente, el objeto de la propuesta y la documentación que integra la presentación.]')
    heading(d,'Objeto y alcance de la contratación')
    heading(d,'Alcance incluido',2);d.add_paragraph('[Describir los bienes y servicios incluidos, las cantidades y los límites de la prestación.]')
    heading(d,'Requisitos y referencias del pliego',2)
    table(d,[['Referencia','Respuesta de la oferta'],['[Cláusula]','[Requisito y ubicación de la respuesta]'],['[Cláusula]','[Requisito y ubicación de la respuesta]']],[32,133])
    heading(d,'Supuestos y exclusiones',2);d.add_paragraph('[Indicar dependencias y prestaciones no incluidas, sin contradecir el pliego.]')
    heading(d,'Solución técnica y plan de trabajo',newpage=True)
    heading(d,'Arquitectura y especificaciones',2);d.add_paragraph('[Describir la solución, sus componentes y las especificaciones ofrecidas.]')
    heading(d,'Requisitos de instalación y puesta en servicio',3);d.add_paragraph('[Detallar las condiciones necesarias y las verificaciones previas a la puesta en servicio.]')
    heading(d,'Entregables y aceptación',2)
    table(d,[['Entregable','Evidencia de aceptación'],['[Entregable 1]','[Documento, prueba o resultado verificable]'],['[Entregable 2]','[Documento, prueba o resultado verificable]']],[62,103])
    heading(d,'Cronograma y responsables',2)
    table(d,[['Etapa','Plazo','Responsable'],['[Etapa 1]','[Plazo]','[Responsable]'],['[Etapa 2]','[Plazo]','[Responsable]']],[74,31,60])
    heading(d,'Oferta económica')
    d.add_paragraph('Moneda: [Completar]    Tratamiento de impuestos: [Completar]')
    table(d,[['Ítem','Concepto','Cant.','Unitario','Importe'],['1','[Servicio o suministro]','[Cant.]','[Precio]','[Importe]'],['2','[Servicio o suministro]','[Cant.]','[Precio]','[Importe]']],[13,64,18,35,35],[WD_ALIGN_PARAGRAPH.CENTER,WD_ALIGN_PARAGRAPH.LEFT,WD_ALIGN_PARAGRAPH.RIGHT,WD_ALIGN_PARAGRAPH.RIGHT,WD_ALIGN_PARAGRAPH.RIGHT])
    p=d.add_paragraph('Subtotal: [Importe]\nImpuestos: [Importe]\nTotal de la oferta: [Moneda e importe]');p.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    d.add_paragraph('[Adjuntar el detalle de la planilla económica y verificar que coincida con los totales de esta propuesta.]')
    heading(d,'Condiciones comerciales y documentación')
    for title,body in [('Validez de la oferta','[Completar el plazo y su fecha de cómputo conforme a la contratación.]'),('Forma de pago y entrega','[Detallar forma de pago, hitos, lugar y plazo de entrega.]'),('Garantía y soporte','[Precisar cobertura, vigencia, canales de atención y exclusiones aplicables.]'),('Documentación y contacto','[Enumerar los anexos exigidos y completar nombre, cargo y datos de contacto del responsable.]')]:
        heading(d,title,2);d.add_paragraph(body)
    if proof:
        # Stress cases use the same live heading styles and numbering definitions.
        for p in d.paragraphs:
            if p.text=='Arquitectura y especificaciones':p.text='Arquitectura de comunicaciones para la integración de edificios y servicios de infraestructura digital'
            elif p.text=='Requisitos de instalación y puesta en servicio':p.text='Requisitos de instalación y puesta en servicio de los equipos de comunicaciones en las sedes incluidas'
        for i in range(6,13):
            heading(d,'Anexo '+str(i-5)+' y documentación complementaria',newpage=i==6)
            d.add_paragraph('Ejemplo de contenido para comprobar la secuencia automática al insertar y mover secciones.')
            if i==10:
                for j in range(1,13):
                    heading(d,'Especificación complementaria '+str(j),2)
                    if j==12:
                        heading(d,'Condiciones de interoperabilidad y validación técnica de los servicios que forman parte de la contratación',3)
                        d.add_paragraph('Esta segunda línea debe conservar la sangría del título y quedar separada del número 10.12.1.')
    return d

def main():
    OUT.mkdir(exist_ok=True)
    save(full(),'oferta-completa')
    d=new_doc('Propuesta comercial','[Nombre del proyecto o servicio]')
    heading(d,'Objetivo y alcance');d.add_paragraph('[Resultado esperado, prestaciones incluidas y criterio de entrega.]')
    heading(d,'Inversión');table(d,[['Concepto','Importe'],['[Servicio o suministro]','[Moneda e importe]'],['Total con impuestos','[Moneda e importe]']],[120,45])
    heading(d,'Condiciones comerciales');d.add_paragraph('[Moneda, impuestos, validez, forma de pago y plazo de entrega.]')
    save(d,'resumen-comercial')
    d=new_doc('Asunto de la comunicación','[Referencia del documento]',letter=True)
    p=d.add_paragraph('De nuestra consideración:');p.paragraph_format.space_before=Pt(22)
    d.add_paragraph('[Presentar el motivo de la comunicación y la información que necesita el destinatario.]')
    d.add_paragraph('[Desarrollar los antecedentes, el alcance o la solicitud. Mantener un párrafo por idea y agregar los anexos que correspondan.]')
    d.add_paragraph('[Indicar el próximo paso y los datos de contacto para continuar la gestión.]')
    p=d.add_paragraph('Atentamente,');p.paragraph_format.space_before=Pt(14)
    p=d.add_paragraph('[Nombre y apellido]');p.paragraph_format.space_before=Pt(28);p.paragraph_format.space_after=Pt(2);p.runs[0].bold=True
    d.add_paragraph('[Cargo]\nULTIMA MILLA S.A.\n[Datos de contacto]');save(d,'membrete')
    save(full(True),'prueba-licitacion',False)
    (ROOT/'especificacion.json').write_text(json.dumps({'font':FONT,'body_pt':12,'heading_pt':SIZES,'number_pt':SIZES,'hanging_tab_mm':TABS,'logo_width_mm':54,'page_mm':[210,297],'margins_mm':[31,20,24,25]},indent=2))

if __name__=='__main__':main()
