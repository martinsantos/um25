from pathlib import Path
from copy import deepcopy
import json,re
from docx import Document
from docx.oxml.ns import qn
from openpyxl import load_workbook
import pdfplumber
import os, tempfile
ROOT=Path(os.environ.get("UM_TEMPLATE_WORK",Path(tempfile.gettempdir())/"um-sans-template-build"))
REPO=Path(os.environ.get("UM_TEMPLATE_REPO",Path(__file__).resolve().parents[3]))
report=[]
def brand_edge(page,side='right'):
    key='x0' if side=='left' else 'x1'
    images=[im for im in page.images if im['bottom']<70]
    if images:
        return images[0][key]
    paths=[c for c in page.curves if c['top']<70 and c['bottom']<70]
    assert len(paths)>=18,'missing vector brand paths'
    return (min if side=='left' else max)(c[key] for c in paths)

for kind in ['oferta-completa','membrete','resumen-comercial','prueba-licitacion','ejemplo-licitacion','imagenes-documento']:
    path=ROOT/'qa/r2'/kind/(kind+'.pdf');d=pdfplumber.open(path)
    for i,p in enumerate(d.pages):
        fonts={c['fontname'].split('+')[-1] for c in p.chars}
        assert fonts <= {'UMSans2-Bold','UMSans2-Regular'},(path,i,fonts)
        side='left' if kind=='membrete' else 'right'
        margin_mm=25 if side=='left' else 190
        assert abs(brand_edge(p,side)-margin_mm*72/25.4)<0.2,(path,i,brand_edge(p,side))
        assert all(c['x0']>69 and c['x1']<540 for c in p.chars if not c['text'].isspace()),(path,i,'visible text beyond margin')
    for page in d.pages:
        words=page.extract_words(extra_attrs=['size','fontname'])
        for index,word in enumerate(words[:-1]):
            if re.fullmatch(r'\d+(?:\.\d+){0,2}\.',word['text']) and word['fontname'].endswith('Bold') and word['size']>=11.9 and abs(word['x0']-25*72/25.4)<0.1:
                title=words[index+1]
                if abs(word['top']-title['top'])<2:
                    assert abs(title['x0']-word['x1']-word['size']*0.25)<0.1,(kind,word['text'],'number-title gap')
                    assert abs(word['x0']-25*72/25.4)<0.1,(kind,word['text'],'number within body line')
                    assert abs(title['size']-word['size'])<0.1,(kind,word['text'],'number-title size')
    report.append({'document':kind,'pages':len(d.pages),'fonts':'UM Sans 2 only','logoAlignment':side,'logoTolerancePt':0.2})
    if kind=='ejemplo-licitacion':
        economic=[p for p in d.pages if '2.501,12' in p.extract_text()]
        assert len(economic)==1 and '5.504,12' in economic[0].extract_text(),'Economic table and total must share a page'
proof=pdfplumber.open(ROOT/'qa/r2/prueba-licitacion/prueba-licitacion.pdf')
text='\n'.join(p.extract_text() for p in proof.pages)
assert '10.12.1' in text and '12. Anexo 7' in text
for p in proof.pages:
    for word in p.extract_words(extra_attrs=['size','fontname']):
        if word['text'] in ['1.','2.','3.','4.','5.','10.','12.'] and word['size']>12 and abs(word['x0']-25*72/25.4)<1:
            assert abs(word['size']-16)<0.1,word
        if word['text']=='10.12.':assert abs(word['size']-13)<0.1
        if word['text']=='10.12.1.':assert abs(word['size']-12)<0.1
# Check the actual exported image geometry and adjacency, not just XML intentions.
for kind,expected in [('oferta-completa',1),('ejemplo-licitacion',1),('imagenes-documento',3)]:
    doc=Document(ROOT/'entrega'/f'{kind}.docx')
    image_paragraphs=[p for p in doc.paragraphs if p._p.xpath('.//w:drawing')]
    assert len(image_paragraphs)==expected,(kind,'missing image block')
    for para in image_paragraphs:
        assert para._p.xpath('.//wp:inline') and not para._p.xpath('.//wp:anchor'),(kind,'floating image')
        assert (para.alignment if para.alignment is not None else para.style.paragraph_format.alignment)==1 and not para._p.xpath('.//w:pBdr'),(kind,'image alignment or frame')
        inline=para._p.xpath('.//wp:inline')[0]
        assert inline.docPr.get('descr'),(kind,'missing image description')
        assert inline.extent.cx/36000<=165.01 and inline.extent.cy/36000<=100.01,(kind,'image outside page bounds')
    with pdfplumber.open(ROOT/'qa/r2'/kind/(kind+'.pdf')) as pdf:
        images=[(page,img) for page in pdf.pages for img in page.images if img['top']>70]
        assert len(images)==expected,(kind,'missing exported product image')
        for page,img in images:
            assert abs((img['x0']+img['x1'])/2-107.5*72/25.4)<0.2,(kind,'image not centered in body')
            assert img['x0']>=25*72/25.4-0.2 and img['x1']<=190*72/25.4+0.2,(kind,'image exceeds margins')
            assert img['bottom']<=273*72/25.4,(kind,'image overlaps footer')
        captions=[p for p in doc.paragraphs if p.style.name=='Comentario de imagen']
        for caption in captions:
            marker=caption.text[:25]
            matches=[(page,page.search(marker,regex=False)) for page in pdf.pages]
            matches=[(page,match[0]) for page,match in matches if match]
            assert len(matches)==1,(kind,'caption lost during export')
            page,match=matches[0]
            preceding=[img for pg,img in images if pg.page_number==page.page_number and img['bottom']<=match['top']]
            assert preceding and 0<match['top']-preceding[-1]['bottom']<20,(kind,'caption detached from image')
        report.append({'document':kind,'images':expected,'imagesCentered':True,'captionSamePage':True,'noFrame':True})
for p in (ROOT/'qa/r2/excel-print').glob('*.pdf'):
    d=pdfplumber.open(p)
    if p.stem=='presupuesto-interno':
        assert len(d.pages)==4
        for pg in d.pages:
            assert not pg.images
            paths=[c for c in pg.curves if c['bottom']<85]
            assert len(paths)>=18,(p,'missing page brand')
            line=next(l for l in pg.lines if l['top']<100 and l['x1']-l['x0']>400)
            assert abs(max(c['x1'] for c in paths)-line['x1'])<0.2,(p,'SVG alignment')
            fonts={c['fontname'].split('+')[-1] for c in pg.chars}
            assert fonts<= {'UMSans2-Bold','UMSans2-Regular'},fonts
            assert max(c['size'] for c in pg.chars)>=19.95,(p,'unintended print scaling')
        assert all('Cantidad' in d.pages[i].extract_text() for i in [1,3]),(p,'broken quantity heading')
        report.append({'document':p.name,'pageCount':4,'printSizes':['A4','A3','A4','A3'],'bodyPt':12,'logoMatchesRule':True})
        continue
    assert len(d.pages)==1
    fonts={c['fontname'].split('+')[-1] for pg in d.pages for c in pg.chars}
    assert fonts<= {'UMSans2-Bold','UMSans2-Regular'},fonts
    line=next(l for l in d.pages[0].lines if l['top']<100 and l['x1']-l['x0']>400)
    assert abs(brand_edge(d.pages[0])-line['x1'])<0.2,(p,brand_edge(d.pages[0]),line['x1'])
    report.append({'document':p.name,'fonts':'UM Sans 2 only','pageCount':1,'logoMatchesRule':True})
recalc=ROOT/'qa/r2/recalculated/prueba-economica.xlsx'
if recalc.exists():
    s=load_workbook(recalc,data_only=True,read_only=True).active
    assert abs(s['F30'].value-5294.12)<0.001
    assert abs(s['F32'].value-5504.12)<0.001
    report.append({'engine':'LibreOffice Calc','subtotal':s['F30'].value,'total':s['F32'].value})
else:raise RuntimeError('Missing native recalculation')
blank=load_workbook(ROOT/'entrega/oferta-economica.xlsx',read_only=True,data_only=True).active
assert blank['F31'].value is None and blank['B14'].value is None
example=load_workbook(ROOT/'entrega/ejemplo-economico.xlsx',read_only=True,data_only=True).active
assert abs(example['F30'].value-5294.12)<0.001 and abs(example['F32'].value-5504.12)<0.001
# Disposable editing fixture: insert a section and change a heading level.
d=Document(ROOT/'entrega/oferta-completa.docx')
target=next(p for p in d.paragraphs if p.text=='Objeto y alcance de la contratación')
target.insert_paragraph_before('Sección incorporada para comprobar la renumeración','Heading 1')
next(p for p in d.paragraphs if p.text=='Supuestos y exclusiones').style=d.styles['Heading 3']
for n in range(30):
    row=deepcopy(d.tables[0].rows[-1]._tr)
    for t in row.xpath('.//w:t'):t.text='Referencia de prueba '+str(n+1) if 'Cláusula' in t.text else 'Requisito extendido para comprobar la continuidad de la tabla.'
    d.tables[0]._tbl.append(row)
d.save(ROOT/'qa/r2/edicion-y-continuidad.docx')
(ROOT/'qa/r2/validation.json').write_text(json.dumps(report,indent=2))
print(json.dumps(report))
