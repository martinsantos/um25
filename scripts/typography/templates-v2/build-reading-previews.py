"""Readable web excerpts generated from the delivered documents, not mock content."""
from pathlib import Path
from lxml import html, etree
from openpyxl import load_workbook
import os, json
from html import escape

REPO=Path(os.environ.get('UM_TEMPLATE_REPO',Path(__file__).resolve().parents[3]))
BASE=REPO/'public/downloads/plantillas-um-sans/2026.10.10-r4'
result={}
for kind in ['oferta-completa','resumen-comercial','membrete','ejemplo-licitacion']:
    doc=html.parse(str(BASE/(kind+'.html')))
    main=doc.xpath('//main')[0]
    for node in main.iter():
        node.attrib.pop('contenteditable',None)
        if node.tag=='h1':node.tag='h3';node.set('class','doc-title')
        elif node.tag in ['h2','h3','h4']:
            level=int(node.tag[1])-1;node.tag='h'+str(level+3);node.set('class',f'doc-heading doc-level-{level}')
        if node.get('class')=='table-wrap':
            node.set('tabindex','0');node.set('role','region');node.set('aria-label','Tabla del documento; desplazamiento horizontal disponible')
    result[kind]=''.join(etree.tostring(c,encoding='unicode',method='html') for c in main)

# Read native workbook values, including exported formula results in the example.
for kind in ['oferta-economica','ejemplo-economico']:
    sheet=load_workbook(BASE/(kind+'.xlsx'),data_only=True,read_only=True).active
    sample=kind=='ejemplo-economico'
    def v(cell):
        value=sheet[cell].value
        if value is None:return ''
        if hasattr(value,'strftime'):value=value.strftime('%d/%m/%Y')
        elif isinstance(value,(int,float)) and cell.startswith(('E','F')):value=f'{value:,.2f}'.replace(',','_').replace('.',',').replace('_','.')
        return escape(str(value))
    rows=''.join('<tr>'+''.join('<td>'+v(f'{col}{row}')+'</td>' for col in 'ABCDEF')+'</tr>' for row in range(14,29) if not sample or sheet[f'B{row}'].value is not None)
    result[kind]=f'''<p class="meta">{v('A2')}</p><h3 class="doc-title">{v('A5')}</h3>
<dl class="doc-sheet-meta"><div><dt>Cliente</dt><dd>{v('A8')}</dd></div><div><dt>Licitación o expediente</dt><dd>{v('C8')}</dd></div><div><dt>Proyecto</dt><dd>{v('A11')}</dd></div><div><dt>Fecha · Moneda</dt><dd>{(v('C11')+' · '+v('E11')) if sample else '—'}</dd></div></dl>
<p class="doc-table-hint">Deslizá la tabla para ver todos los importes →</p>
<div class="table-wrap" tabindex="0" role="region" aria-label="Planilla económica; desplazamiento horizontal disponible"><table class="doc-sheet"><colgroup><col style="width:54px"><col><col style="width:84px"><col style="width:64px"><col style="width:160px"><col style="width:160px"></colgroup><thead><tr>{''.join('<th>'+v(f'{col}13')+'</th>' for col in 'ABCDEF')}</tr></thead><tbody>{rows}</tbody></table></div>
<div class="doc-sheet-totals"><p>Subtotal <span>{v('F30') or '—'}</span></p><p>Impuestos <span>{v('F31') or '—'}</span></p><p>Total de la oferta <span>{v('F32') or '—'}</span></p></div>
<p class="meta">{v('A30')}</p><h4 class="doc-sheet-conditions">{v('A34')}</h4><p>{v('A35')}</p><p>{v('A36')}</p>'''
book=load_workbook(BASE/'presupuesto-interno.xlsx',data_only=True,read_only=True)
sheet=book['Resultado']
def cv(c):
    value=sheet[c].value
    if value is None or value=='':return '—'
    if isinstance(value,(int,float)):value=f'{value:,.2f}'.replace(',','_').replace('.',',').replace('_','.')
    return escape(str(value))
body=''.join(f'<tr><td>{cv("A"+str(r))}</td><td>{cv("D"+str(r))}</td><td>{cv("F"+str(r))}</td></tr>' for r in range(14,18))
totals=''.join(f'<p>{cv("A"+str(r))}<span>{cv("F"+str(r))}</span></p>' for r in [20,21,22,23,24,25,27,28])
result['presupuesto-interno']=f'''<h3 class="doc-title">Resultado del presupuesto</h3><p class="meta">Libro vacío para uso interno · cuatro hojas de trabajo</p><div class="table-wrap" tabindex="0" role="region" aria-label="Resultado del presupuesto"><table><thead><tr><th>Tipo de costo</th><th>Costo USD</th><th>Venta USD</th></tr></thead><tbody>{body}</tbody></table></div><div class="doc-sheet-totals">{totals}</div><h4 class="doc-sheet-conditions">Cómo completar el libro</h4><p>En Parámetros se definen cotización, tarifas, márgenes, ajustes y datos de la oferta. Costeo reúne los componentes por renglón o código y sector o piso. Proveedores permite comparar dos precios y elegir una alternativa.</p><p>Los campos están vacíos para completar cada licitación. El resultado se calcula a partir del detalle; los márgenes se aplican sobre la venta. El costo interno y la oferta al cliente requieren documentos diferentes.</p>'''
book.close()
out=REPO/'src/data/brand-template-reading.json';out.write_text(json.dumps(result,ensure_ascii=False,indent=2)+'\n')
print('Seven readable previews generated from delivered HTML and XLSX.')
