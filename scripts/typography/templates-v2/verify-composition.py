"""Render adversarial inputs in the actual template; keep proofs outside delivery."""
from pathlib import Path
from copy import deepcopy
import importlib.util,json,os,subprocess,tempfile
from docx import Document
from docx.table import Table
from docx.oxml.ns import qn
import pdfplumber

ROOT=Path(os.environ.get('UM_TEMPLATE_WORK',Path(tempfile.gettempdir())/'um-sans-template-build'))
REPO=Path(os.environ['UM_TEMPLATE_REPO'])
RUNTIME=Path(os.environ.get('UM_ARTIFACT_RUNTIME',Path.home()/'.cache/codex-runtimes/codex-primary-runtime'))
proof=ROOT/'qa/composition';proof.mkdir(parents=True,exist_ok=True)
spec=importlib.util.spec_from_file_location('templates',REPO/'scripts/typography/templates-v2/build-documents.py')
builder=importlib.util.module_from_spec(spec);spec.loader.exec_module(builder)
doc=Document(ROOT/'entrega/oferta-completa.docx')
values={
 'Destinatario':'Organismo ficticio de Coordinación de Infraestructura Digital para la Integración Territorial y la Conectividad de las Sedes Provinciales',
 'Proyecto':'Provisión e instalación de comunicaciones de acceso y distribución para la integración de edificios públicos y los centros de atención territorial del proyecto ficticio',
 'Fecha':'10/10/2026'
}
metadata=[deepcopy(p._p) for p in doc.paragraphs if p.text.split('\t')[0] in values]
economic=deepcopy(doc.tables[-1]._tbl)
for node in list(doc.element.body):
 if node.tag!=qn('w:sectPr'):doc.element.body.remove(node)
doc.add_paragraph('CONTROL INTERNO · DATOS FICTICIOS',style='Caption')
doc.add_paragraph('Control de composición de licitaciones',style='Title')
for node in metadata:doc.element.body.insert(len(doc.element.body)-1,node)
for p in doc.paragraphs:
 label=p.text.split('\t')[0]
 if label in values:p.runs[-1].text=values[label]
builder.heading(doc,'Cifras y alineación de los datos económicos')
doc.element.body.insert(len(doc.element.body)-1,economic)
table=Table(economic,doc)
table.add_row()
table._tbl.remove(table.rows[-1]._tr)
table._tbl.append(deepcopy(table.rows[-1]._tr))
numbers=[['1','Cableado y puesta en servicio del sistema de comunicaciones','123.456,78','9.876.543.210,12','123.456.789.012,34'],
 ['2','Documentación y verificaciones de aceptación de los enlaces','999.999,99','123.456.789,12','987.654.321.012,34'],
 ['3','Comprobación de referencias e inventario del proyecto','10.000,00','1.234.567,89','12.345.678.900,12']]
for row,entries in zip(table.rows[1:],numbers):
 for cell,value in zip(row.cells,entries):cell.paragraphs[0].runs[0].text=value
normal='La documentación de la instalación identifica los enlaces y los equipos de cada edificio. El relevamiento confirma las condiciones de acceso y los responsables que acompañarán la puesta en servicio.'
builder.heading(doc,'Continuidad de la tabla',newpage=True)
for _ in range(10):doc.add_paragraph(normal)
builder.heading(doc,'Verificación de entregables',2)
builder.table(doc,[['Renglón de prueba','Evidencia de aceptación'],['PRIMERA FILA',normal+' '+normal],['SEGUNDA FILA','Resultado registrado y documento de conformidad.']],[50,115])
builder.heading(doc,'Imagen vertical y comentario',newpage=True)
svg='<svg xmlns="http://www.w3.org/2000/svg" width="200" height="500"><g fill="none" stroke="#333"><rect x="30" y="30" width="140" height="440"/>'
svg+=''.join(f'<rect x="42" y="{48+i*33}" width="116" height="23"/>' for i in range(12))+'</g></svg>'
(proof/'portrait.svg').write_text(svg)
node=RUNTIME/'dependencies/node/bin/node'
script='import {createRequire} from "node:module";const require=createRequire('+json.dumps(str(RUNTIME/'dependencies/node/package.json'))+');await require("sharp")('+json.dumps(str(proof/'portrait.svg'))+',{density:216}).png().toFile('+json.dumps(str(proof/'portrait.png'))+');'
subprocess.run([str(node),'--input-type=module','-e',script],check=True)
builder.figure(doc,proof/'portrait.png','Esquema técnico ficticio vertical','Esquema de prueba — El gráfico comprueba el centrado, el fondo transparente y el comentario. El texto siguiente mantiene una separación mayor para distinguir el cambio de contenido.',width_mm=108,max_height_mm=100)
builder.heading(doc,'Continuidad después del comentario',2);doc.add_paragraph(normal)
target=proof/'composicion.docx';doc.save(target)
renderer=RUNTIME/'plugins/openai-primary-runtime/plugins/documents/skills/documents/render_docx.py'
env={**os.environ,'FONTCONFIG_FILE':str(ROOT/'fontconfig.xml')}
subprocess.run([str(RUNTIME/'dependencies/python/bin/python3'),str(renderer),str(target),'--output_dir',str(proof),'--emit_pdf'],env=env,check=True)
with pdfplumber.open(proof/'composicion.pdf') as pdf:
 text=' '.join(' '.join(page.extract_text().split()) for page in pdf.pages)
 for value in values.values():assert value in text,('missing metadata',value)
 page=pdf.pages[0];words=page.extract_words();starts={label:page.search(label,regex=False)[0]['top'] for label in values}
 for label,next_label in [('Destinatario','Proyecto'),('Proyecto','Fecha')]:
  selected=[word for word in words if starts[label]-1<=word['top']<starts[next_label]-1 and word['text']!=label]
  assert selected and all(word['x0']>=55*72/25.4-.2 for word in selected),(label,'misaligned continuation')
 cells=page.find_tables()[0].extract()
 for row,expected in zip(cells[1:],numbers):
  assert row[0]==expected[0]
  for actual,value in zip(row[2:],expected[2:]):assert actual==value,('split numeric value',actual,value)
 for page in pdf.pages:
  if any(marker in (page.extract_text() or '') for marker in ['PRIMERA FILA','SEGUNDA FILA']):
   assert 'Renglón de prueba' in page.extract_text(),'table continuation missing header'
 images=[(page,image) for page in pdf.pages for image in page.images if image['top']>100]
 assert len(images)==1
 page,image=images[0];assert abs((image['x0']+image['x1'])/2-107.5*72/25.4)<.2
 assert image['bottom']-image['top']<=100*72/25.4+.2
 caption=page.search('Esquema de prueba',regex=False)[0]
 assert 0<caption['top']-image['bottom']<20
 report={'pages':len(pdf.pages),'longMetadataAligned':True,'numbersComplete':True,'tableHeadersRepeated':True,'portraitCentered':True,'captionSamePage':True}
(proof/'validation.json').write_text(json.dumps(report,indent=2))
print(json.dumps(report))
